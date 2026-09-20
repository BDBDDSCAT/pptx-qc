"""Generate a deliberately broken deck so the README sample output is reproducible.

    python examples/make_demo.py            # writes examples/demo_bad.pptx
    pptx-qc examples/demo_bad.pptx
"""

from __future__ import annotations

import sys
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.util import Inches, Pt

HERE = Path(__file__).resolve().parent
OUT = HERE / "demo_bad.pptx"


def tiny_image(path: Path) -> None:
    """A 48x48 PNG — fine as a thumbnail, useless when blown up."""
    img = Image.new("RGB", (48, 48), (120, 160, 200))
    img.save(path)


def build() -> Path:
    prs = Presentation()
    tiny = HERE / "_tiny.png"
    tiny_image(tiny)

    # --- slide 1: dummy text, Simplified-Chinese leak, Meiryo ---------------
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text_frame.text = "株式会社サンプル　提案資料"
    tf = slide.placeholders[1].text_frame

    run = tf.paragraphs[0].add_run()
    run.text = "事業計画のサマリーです。ダミーテキストなので差し替えてください。"
    run.font.name = "Meiryo"
    run.font.size = Pt(18)

    r2 = tf.add_paragraph().add_run()
    r2.text = "实現に向けた取り組みを、段階的に実施していきます。"  # 实 := simplified-only
    r2.font.name = "Meiryo"
    r2.font.size = Pt(16)

    # --- slide 2: Arial on Japanese text, full-width digits, tiny font ------
    slide2 = prs.slides.add_slide(prs.slide_layouts[1])
    tf2 = slide2.placeholders[1].text_frame
    r3 = tf2.paragraphs[0].add_run()
    r3.text = "売上は増加、コストは減少, 利益は２桁成長"
    r3.font.name = "Arial"
    r3.font.size = Pt(16)

    r4 = tf2.add_paragraph().add_run()
    r4.text = "担当: ﾃｽﾄﾀﾛｳ（ABC123）"
    r4.font.size = Pt(8)

    # --- slide 3: overflow, auto-shrink, low-res image, off-slide shape -----
    slide3 = prs.slides.add_slide(prs.slide_layouts[6])  # blank
    box = slide3.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(1.0), Inches(0.3))
    tf3 = box.text_frame
    tf3.word_wrap = True
    tf3.text = (
        "今期の重点施策として、既存顧客の深耕と新規チャネルの開拓を同時に進め、"
        "四半期ごとのKPIで効果を検証しながら、組織全体の実行力を高めてまいります。"
    )

    box2 = slide3.shapes.add_textbox(Inches(3.0), Inches(2.0), Inches(4.0), Inches(1.0))
    tf4 = box2.text_frame
    tf4.text = "自動縮小が効いているボックスの例です。"
    tf4.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE

    slide3.shapes.add_picture(str(tiny), Inches(3.0), Inches(3.5), width=Inches(5.0))

    off = slide3.shapes.add_textbox(Inches(11.5), Inches(0.2), Inches(2.0), Inches(0.5))
    off.text_frame.text = "スライド外にはみ出した図形"

    # --- slide 4: empty title placeholder (kept blank on purpose) -----------
    slide4 = prs.slides.add_slide(prs.slide_layouts[1])
    slide4.placeholders[1].text_frame.text = "本文だけ入れてタイトルを空にしたページ"

    prs.save(OUT)
    tiny.unlink(missing_ok=True)
    return OUT


if __name__ == "__main__":
    path = build()
    print(f"wrote {path}", file=sys.stderr)
