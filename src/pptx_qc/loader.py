"""Presentation loading and normalised iteration over slides, shapes, runs and media."""

from __future__ import annotations

import hashlib
from collections.abc import Iterator
from dataclasses import dataclass
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.oxml.ns import qn

A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"

#: EMU per point
EMU_PER_PT = 12700
#: EMU per inch
EMU_PER_INCH = 914400


@dataclass
class RunInfo:
    slide: int
    shape: str
    text: str
    font: str | None
    size_pt: float | None
    bold: bool = False

    @property
    def is_empty(self) -> bool:
        return not self.text.strip()


@dataclass
class FrameInfo:
    slide: int
    shape: str
    frame: object  # pptx.text.text.TextFrame
    shape_obj: object  # pptx.shapes.base.BaseShape
    placeholder_type: object | None = None
    is_title: bool = False
    is_table_cell: bool = False


@dataclass
class MediaInfo:
    partname: str
    content_type: str
    size: int
    sha1: str
    width_px: int | None = None
    height_px: int | None = None
    used: bool = False
    used_on: list[str] = None  # type: ignore[assignment]
    #: how many separate parts share these bytes
    part_count: int = 1

    def __post_init__(self) -> None:
        if self.used_on is None:
            self.used_on = []


class Doc:
    """A loaded .pptx with convenience iteration helpers."""

    def __init__(self, path: str | Path) -> None:
        self.path = Path(path)
        self.prs = Presentation(str(self.path))
        self._theme_cache: list[tuple[str | None, str | None]] | None = None
        self._media_cache: list[MediaInfo] | None = None

    # -- basics ---------------------------------------------------------------

    @property
    def slide_width(self) -> int:
        return int(self.prs.slide_width or 0)

    @property
    def slide_height(self) -> int:
        return int(self.prs.slide_height or 0)

    def __iter__(self) -> Iterator[tuple[int, object]]:
        yield from enumerate(self.prs.slides, start=1)

    # -- shapes ---------------------------------------------------------------

    @staticmethod
    def walk_shapes(shapes) -> Iterator[object]:
        for shape in shapes:
            try:
                is_group = shape.shape_type == MSO_SHAPE_TYPE.GROUP
            except (AttributeError, ValueError):
                is_group = False
            if is_group:
                yield from Doc.walk_shapes(shape.shapes)
            else:
                yield shape

    def frames(self) -> Iterator[FrameInfo]:
        """Every text-bearing shape in the deck (including table cells)."""
        for index, slide in self.__iter__():
            for shape in self.walk_shapes(slide.shapes):
                if getattr(shape, "has_text_frame", False):
                    ph_type = None
                    try:
                        placeholder = shape.placeholder_format
                        ph_type = placeholder.type
                    except (AttributeError, ValueError):
                        ph_type = None
                    is_title = False
                    try:
                        is_title = bool(shape == slide.shapes.title)
                    except (AttributeError, ValueError):
                        is_title = False
                    yield FrameInfo(
                        slide=index,
                        shape=shape.name,
                        frame=shape.text_frame,
                        shape_obj=shape,
                        placeholder_type=ph_type,
                        is_title=is_title,
                    )
                if getattr(shape, "has_table", False):
                    for row_i, row in enumerate(shape.table.rows):
                        for col_i, cell in enumerate(row.cells):
                            yield FrameInfo(
                                slide=index,
                                shape=f"{shape.name} [cell {row_i},{col_i}]",
                                frame=cell.text_frame,
                                shape_obj=cell,
                                is_table_cell=True,
                            )

    def runs(self) -> Iterator[RunInfo]:
        for info in self.frames():
            for para in info.frame.paragraphs:
                for run in para.runs:
                    yield RunInfo(
                        slide=info.slide,
                        shape=info.shape,
                        text=run.text,
                        font=run.font.name or para.font.name,
                        size_pt=run.font.size.pt if run.font.size is not None else None,
                        bold=bool(run.font.bold) if run.font.bold is not None else False,
                    )

    def titles(self) -> set[int]:
        """Slide numbers that expose a title placeholder."""
        out: set[int] = set()
        for index, slide in self.__iter__():
            try:
                if slide.shapes.title is not None:
                    out.add(index)
            except (AttributeError, ValueError):
                continue
        return out

    # -- fonts / theme --------------------------------------------------------

    def embedded_fonts(self) -> list[str]:
        """Typefaces declared in <p:embeddedFontLst>."""
        found: list[str] = []
        for node in self.prs.element.iter(qn("p:embeddedFontLst")):
            for font_node in node.iter(qn("p:embeddedFont")):
                tf = font_node.find(qn("p:typeface"))
                if tf is not None and tf.text:
                    found.append(tf.text)
        return found

    def theme_fonts(self) -> list[tuple[str | None, str | None]]:
        """(majorLatin, minorLatin) per slide master."""
        if self._theme_cache is not None:
            return self._theme_cache
        out: list[tuple[str | None, str | None]] = []
        for master in self.prs.slide_masters:
            major = minor = None
            try:
                from pptx.opc.constants import RELATIONSHIP_TYPE as RT

                part = master.part.part_related_by(RT.THEME)
                root = etree.fromstring(part.blob)
                node = root.find(f".//{{{A_NS}}}majorFont/{{{A_NS}}}latin")
                if node is not None:
                    major = node.get("typeface")
                node = root.find(f".//{{{A_NS}}}minorFont/{{{A_NS}}}latin")
                if node is not None:
                    minor = node.get("typeface")
            except Exception:
                pass
            out.append((major, minor))
        self._theme_cache = out
        return out

    def theme_minor_font(self) -> str | None:
        for _, minor in self.theme_fonts():
            if minor:
                return minor
        return None

    # -- media ----------------------------------------------------------------

    def _blip_partnames(self) -> set[str]:
        """Image partnames referenced by any blip in slides, layouts or masters."""
        names: set[str] = set()
        owners: list[object] = list(self.prs.slides)
        for master in self.prs.slide_masters:
            owners.append(master)
            owners.extend(master.slide_layouts)
        for owner in owners:
            try:
                element = owner.element  # type: ignore[attr-defined]
                part = owner.part  # type: ignore[attr-defined]
            except AttributeError:
                continue
            for blip in element.iter(qn("a:blip")):
                rid = blip.get(qn("r:embed")) or blip.get(qn("r:link"))
                if not rid:
                    continue
                try:
                    image_part = part.related_part(rid)
                except Exception:
                    continue
                names.add(str(image_part.partname))
        return names

    def media(self) -> list[MediaInfo]:
        if self._media_cache is not None:
            return self._media_cache

        referenced = self._blip_partnames()
        by_sha: dict[str, MediaInfo] = {}
        for part in self.prs.part.package.iter_parts():
            ctype = getattr(part, "content_type", "") or ""
            if not ctype.startswith("image/"):
                continue
            partname = str(part.partname)
            if partname.startswith("/docProps") or "/docProps/" in partname:
                continue  # the embedded preview thumbnail, not part of the deck
            blob = part.blob
            sha = hashlib.sha1(blob).hexdigest()
            info = by_sha.get(sha)
            if info is None:
                width = height = None
                try:
                    from io import BytesIO

                    from PIL import Image

                    with Image.open(BytesIO(blob)) as im:
                        width, height = im.size
                except Exception:
                    pass
                info = MediaInfo(
                    partname=str(part.partname),
                    content_type=ctype,
                    size=len(blob),
                    sha1=sha,
                    width_px=width,
                    height_px=height,
                )
                by_sha[sha] = info
            else:
                info.part_count += 1
            if str(part.partname) in referenced:
                info.used = True

        # record human-readable locations from picture shapes
        for index, slide in self.__iter__():
            for shape in self.walk_shapes(slide.shapes):
                if getattr(shape, "shape_type", None) == MSO_SHAPE_TYPE.PICTURE:
                    try:
                        blob = shape.image.blob
                    except Exception:
                        continue
                    sha = hashlib.sha1(blob).hexdigest()
                    info = by_sha.get(sha)
                    if info is not None:
                        info.used = True
                        label = f"slide {index}:{shape.name}"
                        if label not in info.used_on:
                            info.used_on.append(label)

        self._media_cache = list(by_sha.values())
        return self._media_cache
