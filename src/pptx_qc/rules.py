"""The checks themselves.

Rule ids are stable and meant to be quoted in issues and config files:

* ``PPTX0xx`` content / text
* ``PPTX1xx`` fonts
* ``PPTX2xx`` layout geometry
* ``PPTX4xx`` embedded media
* ``PPTX5xx`` typography
"""

from __future__ import annotations

import math
from collections import Counter
from collections.abc import Iterator

from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER
from pptx.oxml.ns import qn

from .config import LATIN_ONLY_FONTS, SAFE_FONTS, WINDOWS_ONLY_FONTS, Settings
from .findings import Finding, Severity, rule
from .loader import EMU_PER_INCH, EMU_PER_PT, Doc
from .textutil import (
    find_simplified_chars,
    has_cjk,
    has_halfwidth_katakana,
    is_fullwidth_alnum,
    text_width_em,
)

BLOCK_PLACEHOLDERS = {
    PP_PLACEHOLDER.TITLE,
    PP_PLACEHOLDER.CENTER_TITLE,
    PP_PLACEHOLDER.SUBTITLE,
    PP_PLACEHOLDER.BODY,
    PP_PLACEHOLDER.OBJECT,
}

#: PowerPoint's default text-box insets, in EMU.
DEFAULT_L_INS = 91440
DEFAULT_T_INS = 45720


def _shape_label(shape) -> str:  # noqa: ANN001
    return getattr(shape, "name", "shape")


# ---------------------------------------------------------------------------
# content
# ---------------------------------------------------------------------------


@rule(
    "PPTX001",
    "ダミーテキスト・仮置き文字の残り",
    Severity.ERROR,
    "Placeholder text that was never replaced ('ダミー', 'TBD', 'サンプル株式会社', 'lorem ipsum', …).",
)
def check_dummy_text(doc: Doc, settings: Settings) -> Iterator[Finding]:
    patterns = settings.all_dummy_patterns()
    for frame in doc.frames():
        text = frame.frame.text
        if not text.strip():
            continue
        haystack = text.lower()
        hits: list[str] = []
        for pattern in patterns:
            if pattern.lower() in haystack and pattern not in hits:
                hits.append(pattern)
            if len(hits) >= 3:
                break
        for pattern in hits:
            yield Finding(
                rule="PPTX001",
                severity=Severity.ERROR,
                message=f"ダミーテキストの疑い: 「{pattern}」",
                slide=frame.slide,
                shape=frame.shape,
                hint="顧客提供の実データに差し替えるか、[データ要確認] として明示してください。",
            )


@rule(
    "PPTX002",
    "空のプレースホルダ",
    Severity.WARNING,
    "Title/body placeholders left empty — they show 'Click to edit' in edit mode and break outline view.",
)
def check_empty_placeholders(doc: Doc, settings: Settings) -> Iterator[Finding]:
    for frame in doc.frames():
        if frame.placeholder_type not in BLOCK_PLACEHOLDERS:
            continue
        if frame.frame.text.strip():
            continue
        yield Finding(
            rule="PPTX002",
            severity=Severity.WARNING,
            message="空のプレースホルダが残っています",
            slide=frame.slide,
            shape=frame.shape,
            hint="削除するか内容を入れてください。編集モードで「クリックして編集」が表示されます。",
        )


@rule(
    "PPTX003",
    "スライドタイトルなし",
    Severity.INFO,
    "Slides without a title placeholder hurt outline view, PDF bookmarks and screen readers.",
)
def check_missing_title(doc: Doc, settings: Settings) -> Iterator[Finding]:
    with_title = doc.titles()
    for index, _slide in doc:
        if index not in with_title:
            yield Finding(
                rule="PPTX003",
                severity=Severity.INFO,
                message="タイトルプレースホルダがありません",
                slide=index,
                hint="アウトライン表示・アクセシビリティ・PDF しおりに影響します。",
            )


@rule(
    "PPTX004",
    "簡体字の混入",
    Severity.ERROR,
    "Simplified-Chinese-only characters found in Japanese text — the classic 'Chinese leak' in JP decks.",
)
def check_simplified_chinese(doc: Doc, settings: Settings) -> Iterator[Finding]:
    for run in doc.runs():
        if run.is_empty:
            continue
        for simplified, japanese in find_simplified_chars(run.text):
            yield Finding(
                rule="PPTX004",
                severity=Severity.ERROR,
                message=f"簡体字「{simplified}」が混入（日本語表記は「{japanese}」）",
                slide=run.slide,
                shape=run.shape,
                hint="中国語原稿からのコピペが原因のことが多いです。",
            )


# ---------------------------------------------------------------------------
# fonts
# ---------------------------------------------------------------------------


def _font_usage(doc: Doc) -> Counter:
    counts: Counter = Counter()
    first: dict[str, tuple[int, str]] = {}
    for run in doc.runs():
        if run.is_empty or not run.font:
            continue
        key = run.font.strip()
        counts[key] += 1
        first.setdefault(key, (run.slide, run.shape))
    return counts, first  # type: ignore[return-value]


@rule(
    "PPTX101",
    "Windows 専用フォント（非埋め込み）",
    Severity.WARNING,
    "Fonts such as Meiryo, Yu Gothic or ＭＳ Ｐゴシック that only exist on Windows.",
)
def check_windows_only_fonts(doc: Doc, settings: Settings) -> Iterator[Finding]:
    embedded = {f.lower() for f in doc.embedded_fonts()}
    counts, first = _font_usage(doc)
    for font, count in sorted(counts.items()):
        if font.lower() not in WINDOWS_ONLY_FONTS:
            continue
        if font.lower() in embedded:
            continue
        slide, shape = first[font]
        suffix = f"（他 {count - 1} 箇所）" if count > 1 else ""
        yield Finding(
            rule="PPTX101",
            severity=Severity.WARNING,
            message=f"Windows 専用フォント「{font}」が未埋め込み{suffix}",
            slide=slide,
            shape=shape,
            hint="Mac / Google スライドで書体が変わり、改行位置がずれます。Noto / ヒラギノに置換するか PDF 同梱を。",
        )


@rule(
    "PPTX102",
    "日本語テキストに欧文フォント",
    Severity.WARNING,
    "Japanese glyphs rendered through a Latin-only family — glyph shapes and metrics change per device.",
)
def check_latin_font_on_japanese(doc: Doc, settings: Settings) -> Iterator[Finding]:
    counts, first = _font_usage(doc)
    japanese_runs: Counter = Counter()
    for run in doc.runs():
        if run.font and has_cjk(run.text):
            japanese_runs[run.font.strip()] += 1
    for font, count in sorted(japanese_runs.items()):
        if font.lower() not in LATIN_ONLY_FONTS:
            continue
        slide, shape = first.get(font, (None, None))
        suffix = f"（該当 {count} 箇所）" if count > 1 else ""
        yield Finding(
            rule="PPTX102",
            severity=Severity.WARNING,
            message=f"日本語テキストに欧文フォント「{font}」を指定{suffix}",
            slide=slide,
            shape=shape,
            hint="開いた環境によって字形が変わり、全角文字がはみ出す原因になります。",
        )


@rule(
    "PPTX103",
    "フォント未埋め込み",
    Severity.INFO,
    "Deck uses non-universal fonts but embeds nothing, so rendering depends on the reader's machine.",
)
def check_no_embedded_fonts(doc: Doc, settings: Settings) -> Iterator[Finding]:
    if doc.embedded_fonts():
        return
    counts, _ = _font_usage(doc)
    risky = sorted(f for f in counts if f.lower() not in SAFE_FONTS)
    if not risky:
        return
    preview = "、".join(risky[:6])
    more = f" 他 {len(risky) - 6} 種" if len(risky) > 6 else ""
    yield Finding(
        rule="PPTX103",
        severity=Severity.INFO,
        message=f"フォントが埋め込まれていません（要確認フォント {len(risky)} 種: {preview}{more}）",
        hint="配布前に [ファイル > オプション > 保存] の「フォントを埋め込む」または PDF 化で固定できます。",
    )


@rule(
    "PPTX104",
    "半角カナ",
    Severity.WARNING,
    "Half-width katakana found — should be full-width in Japanese business documents.",
)
def check_halfwidth_katakana(doc: Doc, settings: Settings) -> Iterator[Finding]:
    for run in doc.runs():
        if has_halfwidth_katakana(run.text):
            yield Finding(
                rule="PPTX104",
                severity=Severity.WARNING,
                message=f"半角カナを検出: 「{run.text.strip()[:20]}」",
                slide=run.slide,
                shape=run.shape,
                hint="全角カナに統一してください。",
            )


@rule(
    "PPTX105",
    "使用フォントの種類が多すぎる",
    Severity.INFO,
    "Too many distinct font families make a deck look hand-edited.",
)
def check_font_variety(doc: Doc, settings: Settings) -> Iterator[Finding]:
    counts, _ = _font_usage(doc)
    if len(counts) <= settings.max_font_families:
        return
    names = "、".join(sorted(counts))
    yield Finding(
        rule="PPTX105",
        severity=Severity.INFO,
        message=f"フォントが {len(counts)} 種使われています（推奨 {settings.max_font_families} 種以下）",
        hint=f"内訳: {names}",
    )


# ---------------------------------------------------------------------------
# layout geometry
# ---------------------------------------------------------------------------


def _find(xpath: str, tx_body):  # noqa: ANN001
    """Search a txBody tree for a single node."""
    return tx_body.find(xpath)


def _norm_autofit(tx_body) -> tuple[float | None, float | None]:  # noqa: ANN001
    node = _find(f".//{qn('a:normAutofit')}", tx_body)
    if node is None:
        return None, None
    scale = node.get("fontScale")
    lnspc = node.get("lnSpcReduction")
    return (
        float(scale) / 1000.0 if scale is not None else None,
        float(lnspc) / 1000.0 if lnspc is not None else None,
    )


def _insets(tx_body) -> tuple[int, int, int, int]:
    if tx_body is None:
        return DEFAULT_L_INS, DEFAULT_T_INS, DEFAULT_L_INS, DEFAULT_T_INS
    body_pr = tx_body.find(qn("a:bodyPr"))
    if body_pr is None:
        return DEFAULT_L_INS, DEFAULT_T_INS, DEFAULT_L_INS, DEFAULT_T_INS

    def get(attr: str, default: int) -> int:  # noqa: ANN001
        value = body_pr.get(attr)
        try:
            return int(value) if value is not None else default
        except ValueError:
            return default

    return (
        get("lIns", DEFAULT_L_INS),
        get("tIns", DEFAULT_T_INS),
        get("rIns", DEFAULT_L_INS),
        get("bIns", DEFAULT_T_INS),
    )


def _paragraph_metrics(paragraph, default_size: float = 18.0):  # noqa: ANN001
    text = paragraph.text
    size = default_size
    for run in paragraph.runs:
        if run.font.size is not None:
            size = run.font.size.pt
            break
    spacing = 1.0
    ls = paragraph.line_spacing
    if isinstance(ls, float):
        spacing = ls
    return text, size, spacing


def estimate_overflow(frame_info, settings: Settings):  # noqa: ANN001
    """Return (ratio, required_pt, available_pt) or None when not measurable."""
    shape = frame_info.shape_obj
    width = getattr(shape, "width", None)
    height = getattr(shape, "height", None)
    if not width or not height:
        return None

    tx_body = frame_info.frame._txBody
    l_ins, t_ins, r_ins, b_ins = _insets(tx_body)
    usable_w_pt = max((width - l_ins - r_ins) / EMU_PER_PT, 1.0)
    available_pt = (height - t_ins - b_ins) / EMU_PER_PT

    word_wrap = frame_info.frame.word_wrap
    wrap = True if word_wrap is None else bool(word_wrap)

    scale, lnspc = _norm_autofit(tx_body)
    scale_factor = scale if scale else 1.0
    spacing_reduction = lnspc if lnspc else 0.0

    required_pt = 0.0
    for paragraph in frame_info.frame.paragraphs:
        text, size, spacing = _paragraph_metrics(paragraph)
        size *= scale_factor
        if not text.strip():
            required_pt += size * 1.22
            continue
        em_width = text_width_em(text)
        width_pt = em_width * size
        if wrap:
            lines = max(1, math.ceil(width_pt / usable_w_pt))
        else:
            lines = 1
        line_height = size * 1.22 * spacing * (1.0 - spacing_reduction)
        required_pt += lines * line_height

    if available_pt <= 0:
        return None
    return required_pt / available_pt, required_pt, available_pt


@rule(
    "PPTX201",
    "テキストのはみ出し（推定）",
    Severity.WARNING,
    "Heuristic: estimated text height exceeds the shape, so text likely spills outside the box.",
)
def check_overflow(doc: Doc, settings: Settings) -> Iterator[Finding]:
    for frame in doc.frames():
        if frame.is_table_cell:
            continue
        result = estimate_overflow(frame, settings)
        if result is None:
            continue
        ratio, required, available = result
        tx_body = frame.frame._txBody
        if _find(f".//{qn('a:normAutofit')}", tx_body) is not None:
            continue  # PowerPoint shrinks it; PPTX202 reports that instead
        if ratio <= settings.overflow_tolerance:
            continue
        ratio_label = f"{ratio:.0%}" if ratio < 10 else f"{ratio:.1f}倍"
        yield Finding(
            rule="PPTX201",
            severity=Severity.WARNING,
            message=(
                f"テキストが収まらない可能性: 必要 約{required:.0f}pt / 領域 約{available:.0f}pt"
                f"（{ratio_label}）"
            ),
            slide=frame.slide,
            shape=frame.shape,
            hint="推定値です（文字幅の近似）。実際に開いて確認し、自動調整かフォント縮小を。",
        )


@rule(
    "PPTX202",
    "自動縮小が適用済み",
    Severity.INFO,
    "'Shrink text on overflow' is on, meaning the author already fought the box — size may differ per viewer.",
)
def check_autofit_shrink(doc: Doc, settings: Settings) -> Iterator[Finding]:
    for frame in doc.frames():
        tx_body = frame.frame._txBody
        if _find(f".//{qn('a:normAutofit')}", tx_body) is None:
            continue
        scale, lnspc = _norm_autofit(tx_body)
        if scale is not None and scale >= 0.999 and not lnspc:
            continue
        bits = []
        if scale is not None:
            bits.append(f"文字 {scale:.0%}")
        if lnspc:
            bits.append(f"行間 -{lnspc:.0%}")
        if bits:
            message = "自動縮小が効いています（" + " / ".join(bits) + "）"
        else:
            message = "自動縮小（normAutofit）が有効です"
        yield Finding(
            rule="PPTX202",
            severity=Severity.INFO,
            message=message,
            slide=frame.slide,
            shape=frame.shape,
            hint="意図したサイズなら問題なし。他環境で再計算されると崩れることがあります。",
        )


@rule(
    "PPTX203",
    "図形がスライド外にはみ出し",
    Severity.WARNING,
    "Shapes placed outside the slide area are silently clipped on export.",
)
def check_offslide(doc: Doc, settings: Settings) -> Iterator[Finding]:
    sw, sh = doc.slide_width, doc.slide_height
    if not sw or not sh:
        return
    tol_x = sw * 0.02
    tol_y = sh * 0.02
    for index, slide in doc:
        for shape in Doc.walk_shapes(slide.shapes):
            left, top = getattr(shape, "left", None), getattr(shape, "top", None)
            width, height = getattr(shape, "width", None), getattr(shape, "height", None)
            if None in (left, top, width, height):
                continue
            if (
                left + width < -tol_x
                or top + height < -tol_y
                or left > sw + tol_x
                or top > sh + tol_y
            ):
                yield Finding(
                    rule="PPTX203",
                    severity=Severity.WARNING,
                    message="スライド外（またはほぼ外）に配置されています",
                    slide=index,
                    shape=_shape_label(shape),
                    hint="PDF 出力時に見切れます。意図的な断ち落としなら無視してください。",
                )


@rule(
    "PPTX204",
    "小さすぎる文字",
    Severity.INFO,
    "Runs below the configured minimum font size.",
)
def check_small_font(doc: Doc, settings: Settings) -> Iterator[Finding]:
    if settings.min_font_size <= 0:
        return
    seen: set[tuple[int, str]] = set()
    for run in doc.runs():
        if run.is_empty or run.size_pt is None:
            continue
        if run.size_pt >= settings.min_font_size:
            continue
        key = (run.slide, run.shape)
        if key in seen:
            continue
        seen.add(key)
        yield Finding(
            rule="PPTX204",
            severity=Severity.INFO,
            message=f"{run.size_pt:g}pt の文字（基準 {settings.min_font_size:g}pt 未満）",
            slide=run.slide,
            shape=run.shape,
            hint="提案資料では本文 12pt 以上が目安です。--min-font-size で調整できます。",
        )


@rule(
    "PPTX205",
    "テーマフォントの不統一",
    Severity.INFO,
    "Slide masters disagree on their theme fonts, a common sign of slides pasted from another deck.",
)
def check_theme_consistency(doc: Doc, settings: Settings) -> Iterator[Finding]:
    themes = doc.theme_fonts()
    distinct = {t for t in themes if any(t)}
    if len(distinct) <= 1:
        return
    rendered = " / ".join(f"{major or '-'} + {minor or '-'}" for major, minor in sorted(distinct))
    yield Finding(
        rule="PPTX205",
        severity=Severity.INFO,
        message=f"スライドマスター間でテーマフォントが {len(distinct)} 通りあります",
        hint=f"内訳: {rendered}。別資料からの貼り付けが疑われます。",
    )


# ---------------------------------------------------------------------------
# media
# ---------------------------------------------------------------------------


@rule(
    "PPTX401",
    "画像解像度が低い",
    Severity.WARNING,
    "Picture displayed larger than its pixel data supports (effective dpi below threshold).",
)
def check_low_dpi(doc: Doc, settings: Settings) -> Iterator[Finding]:
    media_by_sha = {m.sha1: m for m in doc.media()}
    worst: dict[str, tuple[float, int, str]] = {}
    for index, slide in doc:
        for shape in Doc.walk_shapes(slide.shapes):
            if getattr(shape, "shape_type", None) != MSO_SHAPE_TYPE.PICTURE:
                continue
            try:
                import hashlib

                sha = hashlib.sha1(shape.image.blob).hexdigest()
            except Exception:
                continue
            info = media_by_sha.get(sha)
            if info is None or not info.width_px or not info.height_px:
                continue
            width_in = (getattr(shape, "width", 0) or 0) / EMU_PER_INCH
            height_in = (getattr(shape, "height", 0) or 0) / EMU_PER_INCH
            if width_in <= 0 or height_in <= 0:
                continue
            dpi = min(info.width_px / width_in, info.height_px / height_in)
            if dpi >= settings.min_image_dpi:
                continue
            prev = worst.get(sha)
            if prev is None or dpi < prev[0]:
                worst[sha] = (dpi, index, _shape_label(shape))
    for dpi, index, label in worst.values():
        yield Finding(
            rule="PPTX401",
            severity=Severity.WARNING,
            message=f"画像の実効解像度が約 {dpi:.0f}dpi（基準 {settings.min_image_dpi:g}dpi 未満）",
            slide=index,
            shape=label,
            hint="拡大表示や印刷でぼやけます。高解像度版に差し替えるか縮小配置してください。",
        )


@rule(
    "PPTX402",
    "同一画像の重複埋め込み",
    Severity.INFO,
    "The same image bytes are stored several times, inflating the file.",
)
def check_duplicate_media(doc: Doc, settings: Settings) -> Iterator[Finding]:
    for info in doc.media():
        if info.part_count > 1:
            mb = info.size * info.part_count / (1024 * 1024)
            yield Finding(
                rule="PPTX402",
                severity=Severity.INFO,
                message=f"同一画像が {info.part_count} 回埋め込まれています（合計 約{mb:.1f}MB）",
                hint=f"{info.partname} など。PowerPoint の「メディアの圧縮」で 1 つにまとまります。",
            )


@rule(
    "PPTX403",
    "未使用の画像",
    Severity.INFO,
    "Image parts that no slide, layout or master references any more.",
)
def check_unused_media(doc: Doc, settings: Settings) -> Iterator[Finding]:
    for info in doc.media():
        if info.used:
            continue
        yield Finding(
            rule="PPTX403",
            severity=Severity.INFO,
            message=f"未使用の画像（{info.size / 1024:.0f}KB）",
            hint=f"{info.partname}。削除するとファイルが軽くなります。",
        )


# ---------------------------------------------------------------------------
# typography
# ---------------------------------------------------------------------------


@rule(
    "PPTX501",
    "全角/半角の混在",
    Severity.INFO,
    "Full-width and half-width alphanumerics mixed inside one block.",
)
def check_width_mixing(doc: Doc, settings: Settings) -> Iterator[Finding]:
    for frame in doc.frames():
        text = frame.frame.text
        if not text.strip():
            continue
        full = any(is_fullwidth_alnum(c) for c in text)
        half = any(c.isascii() and c.isalnum() for c in text)
        double_space = "　　" in text or "  " in text
        problems = []
        if full and half:
            problems.append("全角英数と半角英数が混在")
        elif full:
            problems.append("全角英数を使用")
        if double_space:
            problems.append("連続スペース")
        if not problems:
            continue
        yield Finding(
            rule="PPTX501",
            severity=Severity.INFO,
            message="／".join(problems),
            slide=frame.slide,
            shape=frame.shape,
            hint="数字・英字は半角、区切りは全角スペース 1 つに統一するのが無難です。",
        )


@rule(
    "PPTX502",
    "約物の不統一",
    Severity.INFO,
    "Japanese and Latin punctuation mixed in the same block (、 vs , / 。 vs .).",
)
def check_punctuation_mixing(doc: Doc, settings: Settings) -> Iterator[Finding]:
    for frame in doc.frames():
        text = frame.frame.text
        if not text.strip():
            continue
        problems = []
        if "、" in text and "," in text:
            problems.append("「、」と「,」が混在")
        if "。" in text and ". " in text:
            problems.append("「。」と「.」が混在")
        if "：" in text and ":" in text:
            problems.append("「：」と「:」が混在")
        if not problems:
            continue
        yield Finding(
            rule="PPTX502",
            severity=Severity.INFO,
            message="／".join(problems),
            slide=frame.slide,
            shape=frame.shape,
            hint="日本語資料は全角約物に統一するのが基本です。",
        )
