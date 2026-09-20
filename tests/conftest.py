"""Helpers for building throwaway decks in tests."""

from __future__ import annotations

from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt


def new_deck():
    """A deck with one title+body slide; returns (prs, slide)."""
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    return prs, slide


def set_text(shape, text: str, font: str | None = None, size: float | None = None):
    tf = shape.text_frame
    tf.text = ""
    run = tf.paragraphs[0].add_run()
    run.text = text
    if font:
        run.font.name = font
    if size:
        run.font.size = Pt(size)
    return run


def tiny_png(path, size=(48, 48)):
    Image.new("RGB", size, (120, 160, 200)).save(path)
    return path


def save(prs, path):
    prs.save(path)
    return path


INCH = Inches
